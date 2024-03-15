import streamlit as st
from streamlit_extras.stylable_container import stylable_container
import os 
import re
import requests
from bs4 import BeautifulSoup
import io
from io import StringIO, BytesIO
from tqdm.auto import tqdm
from llama_index.core.agent import ReActAgent
from llama_index.core.llms import ChatMessage, MessageRole
from llama_index.core.memory import ChatMemoryBuffer
import llama_index.core
from llama_index.llms.azure_openai import AzureOpenAI
from llama_index.embeddings.azure_openai import AzureOpenAIEmbedding
from llama_index.core import get_response_synthesizer,  ServiceContext, Document
from llama_index.core.schema import Document
import nest_asyncio
__import__('pysqlite3')
import sys
sys.modules['sqlite3'] = sys.modules.pop('pysqlite3')
sys.path.append('/home/qcells/Desktop/rag_project/utils')
sys.path.append("/workspaces/hanwhaqcells/utils")
sys.path.append('../utils')
from qcells_route_engine import qcell_engine, web_engine, Decide_to_search_external_web, GoogleRandomSearchToolSpec, VectordbSearchToolSpec
from qcells_custom_rag import create_db_chat, pptx_load_data, pdf_load_data, get_youtube_metadata, get_timeline
from web_crack import RemoteDepthReader
from youtube_transcript_api import YouTubeTranscriptApi
from streamlit_pdf_viewer import pdf_viewer
import phoenix as px
from phoenix.trace import DocumentEvaluations, SpanEvaluations
from phoenix.session.evaluation import get_qa_with_reference, get_retrieved_documents
from phoenix.evals import (
    HallucinationEvaluator,
    OpenAIModel,
    QAEvaluator,
    RelevanceEvaluator,
    run_evals,
)
import base64
import os
from pptx import Presentation
import pdfkit

nest_asyncio.apply()
st.set_page_config(page_title="RAG",  layout="wide",  page_icon="⛅")

@st.cache_resource
def lanch_px_app():
    llama_index.core.set_global_handler("arize_phoenix")
    session = px.launch_app(host='0.0.0.0')
    
lanch_px_app()
def set_llm():
    return AzureOpenAI(
            model="gpt-35-turbo",
            deployment_name="qcell_gpt_model",
            temperature = 0,
            api_key="c11ed4df2d35412b89a7b51a631bf0e4",
            azure_endpoint="https://rag-openai-qcells-east.openai.azure.com/",
            api_version="2024-02-15-preview")
def set_rag():
    return AzureOpenAI(
            model="gpt-35-turbo",
            deployment_name="qcell-gpt-model-rag",
            temperature = 0,
            api_key="c11ed4df2d35412b89a7b51a631bf0e4",
            azure_endpoint="https://rag-openai-qcells-east.openai.azure.com/",
            api_version="2024-02-15-preview")
def set_llm4():
    return AzureOpenAI(
            model="gpt-4",
            deployment_name="qcell_gpt4_model",
            temperature = 0,
            api_key="2b6d6cbbc0ae4276aad07db896f63bfd",
            azure_endpoint="https://rag-openai-qcells-norway.openai.azure.com/",
            api_version="2024-02-15-preview")
def set_embedding():
    return AzureOpenAIEmbedding(
        model="text-embedding-ada-002",
        deployment_name="qcell_embedding_model",
        api_key="c11ed4df2d35412b89a7b51a631bf0e4",
        azure_endpoint="https://rag-openai-qcells-east.openai.azure.com/",
     api_version="2023-07-01-preview")
    
def get_answer_yn(llm, query_str, text_chunks):
    synthesizer = get_response_synthesizer(llm = llm, response_mode="refine", output_cls = Decide_to_search_external_web)
    result_response = synthesizer.get_response(query_str = query_str, text_chunks=[text_chunks], verbose = True)      
    return result_response


if "llm" not in st.session_state:
    st.session_state.llm = set_llm()

if "llm_rag" not in st.session_state:
    st.session_state.llm_rag = set_rag()

if "embedding" not in st.session_state:
    st.session_state.embedding = set_embedding()
    
if "llm4" not in st.session_state:
    st.session_state.llm4 = set_llm4()
    
if "rag" not in st.session_state:    
    st.session_state.rag = qcell_engine(llm = st.session_state.llm_rag, embedding = st.session_state.embedding)

if "webrag" not in st.session_state:
    st.session_state.webrag = web_engine(llm = st.session_state.llm_rag, embedding = st.session_state.embedding)

if "service_context" not in st.session_state:
    st.session_state.service_context = ServiceContext.from_defaults(llm=st.session_state.llm,embed_model=st.session_state.embedding,)

if "doc_read_yn" not in st.session_state:
    st.session_state.doc_read_yn = False
if "youtube_read_yn" not in st.session_state:
    st.session_state.youtube_read_yn = False
if "webpage_read_yn" not in st.session_state:
    st.session_state.webpage_read_yn = False
if "webpageall_read_yn" not in st.session_state:
    st.session_state.webpageall_read_yn = False
if "external_docs" not in st.session_state:
    st.session_state.external_docs = []
if "doc_name" not in st.session_state:
    st.session_state.display_datasource = ''
if "chat_db" not in st.session_state:
    st.session_state.chat_db = None
if "messages1" not in st.session_state.keys():
    st.session_state.messages1 = [{"role": "assistant", "content": "Hi"}]
if "messages2" not in st.session_state.keys(): 
    st.session_state.messages2 = [{"role": "assistant", "content": "Hi"}]
if "messages3" not in st.session_state.keys(): 
    st.session_state.messages3 = [{"role": "assistant", "content": "Hi"}]
if "messages4" not in st.session_state.keys(): 
    st.session_state.messages4 = [{"role": "assistant", "content": "Hi"}]
if "chatgpt_mode" not in st.session_state:
    st.session_state.chatgpt_mode = 'ChatGPT 3.5'
if "chat_engine4" not in st.session_state.keys(): 
    st.session_state.session_state = None
if "chat_engine" not in st.session_state.keys(): 
    st.session_state.prompts1 = []
    st.session_state.prompts2 = []
    st.session_state.prompts3 = []
    st.session_state.prompts4 = []

def make_data_instance(_docs):
    st.session_state.messages4 = [{"role": "assistant", "content": "Hi"}]    
    st.session_state.prompts4 = []
    with st.spinner("Creating knowledge database"):
        st.session_state.chat_db = create_db_chat(_docs, st.session_state.llm_rag, st.session_state.embedding,  st.session_state.service_context)
    memory = ChatMemoryBuffer.from_defaults(token_limit=2000)
    st.session_state.chat_engine2 = ReActAgent.from_llm(st.session_state.chat_db.to_tool_list(), memory = memory, llm = st.session_state.llm_rag, verbose = True)    
    st.session_state.model_select = 'ChatGPT+CustomRAG'

with st.sidebar.expander("DOCUMENT"): 
    multiple_files = st.file_uploader("Drag & Drop:", accept_multiple_files=True, key="file_uploader")
    if len(multiple_files) > 0:
        st.session_state.external_docs = []
            
        for doc in multiple_files:
            if doc.name.split('.')[-1] == 'pdf':
                st.session_state.display_datasource = {'pdf': multiple_files[0].getvalue()}
                string_data = pdf_load_data(doc)
                pdf_documents = [Document(text=string_data, metadata= {"title" : doc.name, 'resource' : 'file'})]
                st.session_state.external_docs.append(pdf_documents)  
                
            if doc.name.split('.')[-1] == 'py':
                stringio = StringIO(doc.getvalue().decode("utf-8"))
                string_data = stringio.read()
                document = [Document(text=string_data, metadata= {"title" : doc.name, 'resource' : 'file'})]
                st.session_state.external_docs.append(document)  

            if doc.name.split('.')[-1] == 'pptx':
                result = pptx_load_data(doc) 
                document = [Document(text=result, metadata={"title" : doc.name, 'resource' : 'file'})]
                st.session_state.external_docs.append(document)    
        
    btn_doc = st.button("SAVE", on_click = make_data_instance, args=[st.session_state.external_docs], key="btn_doc")  

with st.sidebar.expander("YOUTUBE"): 
    youtube_data = st.text_input("URL", key='youtubeurl')
    if len(youtube_data) > 0:
        meta_data = get_youtube_metadata(youtube_data)
        st.session_state.display_datasource = {'youtube': meta_data['embed_url']}
        data = YouTubeTranscriptApi.get_transcript(youtube_data.split('v=')[-1])
        documents = []
        for i in data:
            i['div'] = int(i['start'] / 60)  
            
        total_text= []
        distinct_div = set(item['div'] for item in data)
        for idx, d in enumerate(distinct_div):
            texts = []
            k = [i for i in data if i['div'] == d]
            start_time_min = min([i['start'] for i in k])
            start_time_max = max([i['start'] for i in k])
            text_div = [i['text'] for i in k]
            texts.append('[youtube play time] {}\n'.format(get_timeline(start_time_min)) + ' '.join(text_div))
            texts = '\n'.join(texts)
            total_text.append(texts)
        
        total_text = '\n'.join(total_text)
        documents.append(Document(text=total_text, metadata= {'title':  meta_data['title'], 'resource' : 'youtube'}))
        st.session_state.external_docs.append(documents)
    btn_youtube = st.button("SAVE", on_click = make_data_instance, args=[st.session_state.external_docs], key="btn_youtube")  

# with st.sidebar.expander("WEB DEEP SEARCH"): 
#     all_page_data = st.text_input("URL", key = 'all_weburl') 
#     if st.session_state.webpageall_read_yn == False:
#         if len(all_page_data) > 0:
#             web_cracker = RemoteDepthReader()
#             web_crack_documents = web_cracker.load_data(url=all_page_data)
#             web_crack_documents = [i for i in web_crack_documents if i.split('.')[-1] == 'html']
#             with st.spinner("Fetching page texts"):
#                 docs = []
#                 for i in tqdm(web_crack_documents[:50]):
#                     loader = AsyncChromiumLoader([i])
#                     bs_transformer = BeautifulSoupTransformer()
#                     docs_transformed = bs_transformer.transform_documents(loader.load(), unwanted_tags= ['style','script'], tags_to_extract = ['a','p', 'span'])
#                     result_string = remove_nested_parentheses(docs_transformed[0].page_content)
#                     web_data_documents = [Document(text=result_string)]
#                     web_data_documents[0].metadata['resource'] = 'web_allpage'
#                     web_data_documents[0].metadata['title'] = 'web_allpage'
#                     docs.append(web_data_documents[0])
#             st.session_state.external_docs.append(docs)
#             st.session_state.webpageall_read_yn = True
#         else:
#             st.session_state.webpageall_read_yn = False
#     else:
#         pass
#     btn_allpage = st.button("SAVE", on_click = make_data_instance, args=[st.session_state.external_docs], key="btn_allpage")  
    
with st.sidebar.expander("WEB PAGE"): 
    single_page_data = st.text_input("URL", key = 'single_weburl') 
    if len(single_page_data) > 0:
        r=requests.get(single_page_data)
        soup = BeautifulSoup(r.content, 'html.parser')
        anchor_elements = soup.find_all('p')
        result_string = '\n'.join([i.text for i in anchor_elements])                
        web_data_documents = [Document(text=result_string, metadata= {'title':  'web_page', 'resource' : 'web_page'})]
        st.session_state.external_docs.append(web_data_documents)
    btn_singlepage = st.button("SAVE", on_click = make_data_instance, args=[st.session_state.external_docs], key="btn_singlepage")  
    
def convert_message(x):
    return [ChatMessage(role=i['role'], content=i['content']) for i in x]        

def reset_conversation(x):
    st.session_state.llm = set_llm()
    st.session_state.llm4 = set_llm4()
    st.session_state.embedding = set_embedding()
    st.session_state.service_context = ServiceContext.from_defaults(llm=st.session_state.llm,embed_model=st.session_state.embedding,)
    if x == 'ChatGPT 3.5':
        st.session_state.chat_engine = st.session_state.llm
        st.session_state.webrag.reset()
        st.session_state.webrag = web_engine(llm = st.session_state.llm_rag, embedding = st.session_state.embedding)
        st.session_state.messages1 = [{"role": "assistant", "content": "Hi"}]    
        st.session_state.prompts1 = []
    if x == 'ChatGPT 4':
        st.session_state.chat_engine_llm4 = st.session_state.llm4
        st.session_state.webrag.reset()
        st.session_state.webrag = web_engine(llm = st.session_state.llm_rag, embedding = st.session_state.embedding)
        st.session_state.messages2 = [{"role": "assistant", "content": "Hi"}]    
        st.session_state.prompts2 = []
    if x == 'ChatGPT+RAG':
        st.session_state.rag.reset()
        st.session_state.rag = qcell_engine(llm = st.session_state.llm_rag, embedding = st.session_state.embedding)
        st.session_state.messages3 = [{"role": "assistant", "content": "Hi"}]    
        st.session_state.prompts3 = []
    if x == 'ChatGPT+CustomRAG':
        st.session_state.messages4 = [{"role": "assistant", "content": "Hi"}]    
        st.session_state.prompts4 = []

if prompt := st.chat_input("Your question", key = 'chat_input_query'): # Prompt for user input and save to chat history
    if st.session_state.chatgpt_mode == 'ChatGPT 3.5':
        st.session_state.prompts1.append(prompt)
        st.session_state.messages1.append({"role": "user", "content": prompt})
    if st.session_state.chatgpt_mode == 'ChatGPT 4':
        st.session_state.prompts2.append(prompt)
        st.session_state.messages2.append({"role": "user", "content": prompt})
    if st.session_state.chatgpt_mode == 'ChatGPT+RAG':
        st.session_state.prompts3.append(prompt)
        st.session_state.messages3.append({"role": "user", "content": prompt})
    if st.session_state.chatgpt_mode == 'ChatGPT+CustomRAG':
        st.session_state.prompts4.append(prompt)
        st.session_state.messages4.append({"role": "user", "content": prompt})

col1, col2 = st.columns([1, 3])
with col1:
    st.session_state.chosen_id = st.selectbox('Model Name', ('ChatGPT 3.5', 'ChatGPT 4', 'ChatGPT+RAG', 'ChatGPT+CustomRAG'), key = 'model_select')


def chat_box(text):
    # texts = re.findall(r'\```[^\```]+\```|[^\```]+', str(text))
    texts = text.split('```')
    for i in texts:
        with stylable_container(
            "codeblock",
            """
            code {font-size:13px; font-family: Arial; white-space: pre-wrap !important;}
            """,):
            x = st.code(i, language = 'md')    
    return x

if st.session_state.chosen_id == 'None':
    st.session_state.chosen_id = 'ChatGPT 3.5'
    
if st.session_state.chosen_id == "ChatGPT 3.5":
    st.button('Reset Chat', on_click=reset_conversation,args=["ChatGPT 3.5"], key = 'reset1')
    with st.container(height=700):
        st.session_state.chat_engine = st.session_state.llm
        st.session_state.chatgpt_mode = 'ChatGPT 3.5'
        for message in st.session_state.messages1: # Display the prior chat messages
            with st.chat_message(message["role"]):
                msg = message["content"]
                chat_box(msg)
        
        if st.session_state.chatgpt_mode == 'ChatGPT 3.5': 
            if st.session_state.messages1[-1]["role"] == "user":
                with st.chat_message("assistant"):
                    with st.spinner("Thinking..."):
                        prompt_ = convert_message(st.session_state.messages1)
                        response = st.session_state.chat_engine.chat(prompt_) #결과                                            
                        chat_box(response.message.content)
                        message = {"role": "assistant", "content": response.message.content} #저장
                        st.session_state.messages1.append(message) # Add response to message history       
                    answer_eval = get_answer_yn(st.session_state.llm, prompt, response.message.content)
                    res = None
                    if answer_eval.Succeed_answer == False:     
                        if answer_eval.Decide_web_search == True:
                            if answer_eval.Reason == True:
                                with st.spinner("Google + Bing search"):
                                    try:
                                        res = st.session_state.webrag.chat(answer_eval.Searchable_query)
                                        res = res.response
                                    except Exception as e:
                                        chat_engine = ReActAgent.from_tools(GoogleRandomSearchToolSpec().to_tool_list(), max_iterations = 10,llm = st.session_state.llm, verbose = True)
                                        res = 'retry! ' + str(e)
                                        
                                    # if len(res.sources[0].content) > 0:
                                    #     sources = eval(res.sources[0].content)            
                                    #     for source in sources:
                                    #         st.sidebar.write("check out this [link](%s)" % source)            
                if res:
                    with st.chat_message("assistant"):
                        message = {"role": "assistant", "content": res}
                        st.session_state.messages1.append(message) # Add response to message history           
                        chat_box(res)         
                    
if st.session_state.chosen_id == "ChatGPT 4":
    st.button('Reset Chat', on_click=reset_conversation,args=["ChatGPT 4"], key = 'reset2')
    with st.container(height=700):
        st.session_state.chat_engine = st.session_state.llm4
        st.session_state.chatgpt_mode = 'ChatGPT 4'
        for message in st.session_state.messages2: # Display the prior chat messages
            with st.chat_message(message["role"]):
                msg = message["content"]
                chat_box(msg)
        if st.session_state.chatgpt_mode == 'ChatGPT 4': 
            if st.session_state.messages2[-1]["role"] == "user":
                with st.chat_message("assistant"):
                    with st.spinner("Thinking..."):
                        prompt_ = convert_message(st.session_state.messages2)
                        response = st.session_state.chat_engine.chat(prompt_) #결과                    
                        chat_box(response.message.content )                        
                        message = {"role": "assistant", "content": response.message.content} #저장
                        st.session_state.messages2.append(message) # Add response to message history       
                    answer_eval = get_answer_yn(st.session_state.llm,  prompt, response.message.content)
                    res = None
                    if answer_eval.Succeed_answer == False:     
                        if answer_eval.Decide_web_search == True:
                            if answer_eval.Reason == True:
                                with st.spinner("Google + Bing search"):
                                    try:
                                        res = st.session_state.webrag.chat(answer_eval.Searchable_query)
                                        res = res.response
                                    except Exception as e:
                                        chat_engine = ReActAgent.from_tools(GoogleRandomSearchToolSpec().to_tool_list(), max_iterations = 10,llm = st.session_state.llm, verbose = True)
                                        res = 'retry! ' + str(e)
                if res:
                    with st.chat_message("assistant"):
                        message = {"role": "assistant", "content": res}
                        st.session_state.messages1.append(message) # Add response to message history           
                        chat_box(res)         
                            
if st.session_state.chosen_id == "ChatGPT+RAG":  
    st.button('Reset Chat', on_click=reset_conversation,args=["ChatGPT+RAG"], key = 'reset3')   
    with st.container(height=700):
        st.session_state.chat_engine = st.session_state.rag
        st.session_state.chatgpt_mode = 'ChatGPT+RAG'
        for message in st.session_state.messages3: # Display the prior chat messages
            with st.chat_message(message["role"]):
                msg = message["content"]
                chat_box(msg)
        if st.session_state.chatgpt_mode == 'ChatGPT+RAG': 
            if st.session_state.messages3[-1]["role"] == "user":
                with st.chat_message("assistant"):
                    with st.spinner("Thinking..."):
                        try:
                            response = st.session_state.chat_engine.chat(prompt)
                            res = response.response
                        except Exception as e:
                            st.session_state.chat_engine.reset()
                            res = str(e)
                        chat_box(res)
                        message = {"role": "assistant", "content": res}
                        st.session_state.messages3.append(message) # Add response to message history        
    
col1_chat, col2_chat = st.columns([3, 2])
if st.session_state.chosen_id == "ChatGPT+CustomRAG":
    st.sidebar.button('Reset Chat', on_click=reset_conversation,args=['ChatGPT+CustomRAG'], key = 'reset4')    
    with col1_chat.container(height=700):
        st.session_state.chatgpt_mode = 'ChatGPT+CustomRAG'    
        for message in st.session_state.messages4: # Display the prior chat messages
            with st.chat_message(message["role"]):
                msg = message["content"]
                chat_box(msg)
        if st.session_state.chatgpt_mode == 'ChatGPT+CustomRAG':
            try:
                    
                if st.session_state.messages4[-1]["role"] == "user":
                    with st.chat_message("assistant"):
                        with st.spinner("Thinking..."):
                            response = st.session_state.chat_engine2.chat(prompt)
                            res = response.response
                            chat_box(res)
                            message = {"role": "assistant", "content": res}
                            st.session_state.messages4.append(message) # Add response to message history        
                
                if len(st.session_state.chat_db.summary) > 0:
                    for i in st.session_state.chat_db.summary:
                        message = {"role": "assistant", "content": i}
                        st.session_state.messages4.append(message) # Add response to message history        
                        st.session_state.chat_db.summary = []
    
                        with st.chat_message("assistant"):
                            chat_box(i)
                
    
            except Exception as e:
                st.warning('Insert Materials' + str(e), icon="⚠️")
                pass
    with col2_chat.container(height=700):
        if 'pdf' in st.session_state.display_datasource:
            pdf_viewer(input=st.session_state.display_datasource['pdf'], width=550)        
        if 'youtube' in st.session_state.display_datasource:            
            st.write('<iframe src="{}" width="100%" height="400px"></iframe>'.format(st.session_state.display_datasource['youtube']),unsafe_allow_html=True,)


